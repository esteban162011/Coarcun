"""Cuentas Coarcun - lógica del programa (sin interfaz)."""
import difflib
import hashlib
import hmac
import io
import os
import re
import secrets
import sqlite3
import unicodedata
from contextlib import contextmanager
from datetime import date, time

import pandas as pd

DB_PATH = os.environ.get("COARCUN_DB", "cuentas_coarcun.db")
CODIGO_FILE = os.environ.get("COARCUN_CODE_FILE", "codigo_admin.txt")

RESIDENCIAS = ["Fusagasugá", "Bogotá", "La Mesa", "Facatativá",
               "Girardot", "Soacha", "Chía", "Tocancipá"]
ROL_CENTRAL, ROL_ASISTENTE = "Arbitro", "Asistente"
ROLES_COBRABLES = (ROL_CENTRAL, ROL_ASISTENTE)
TODOS = "Todos los torneos"

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


# ---------------------------------------------------------------- utilidades
def norm(s) -> str:
    """Minúsculas, sin tildes y sin espacios repetidos (para comparar nombres)."""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def limpiar(s) -> str:
    return re.sub(r"\s+", " ", str(s)).strip()


def fmt(n) -> str:
    return f"{n:,.0f}".replace(",", ".")


def fecha_bonita(iso: str) -> str:
    d = date.fromisoformat(iso)
    return f"{d.day:02d} de {MESES[d.month - 1].capitalize()}"


def parse_fecha(v):
    """Acepta fechas reales o textos como 'jueves, octubre 15, 2026'."""
    if hasattr(v, "date") and not isinstance(v, str):
        return v.date().isoformat()
    m = re.search(r"([A-Za-zÁÉÍÓÚáéíóúñÑ]+)\s+(\d{1,2}),?\s+(\d{4})", str(v))
    if m and norm(m.group(1)) in [norm(x) for x in MESES]:
        mes = [norm(x) for x in MESES].index(norm(m.group(1))) + 1
        return date(int(m.group(3)), mes, int(m.group(2))).isoformat()
    try:
        return pd.to_datetime(str(v), dayfirst=True).date().isoformat()
    except Exception:
        return None


def parse_hora(v) -> str:
    """Redondea al minuto (el Excel trae basura como 10:59:59.971)."""
    if isinstance(v, time):
        seg = v.hour * 3600 + v.minute * 60 + v.second + v.microsecond / 1e6
        mins = round(seg / 60)
        return f"{(mins // 60) % 24:02d}:{mins % 60:02d}"
    return "" if v is None or pd.isna(v) else str(v)


def normalizar_rol(v) -> str:
    n = norm(v)
    return {"arbitro": ROL_CENTRAL, "asistente": ROL_ASISTENTE}.get(n, limpiar(v).capitalize())


def normalizar_torneo(v) -> str:
    return limpiar(v).rstrip(". ")


# ------------------------------------------------------------- base de datos
@contextmanager
def db():
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    c.execute("PRAGMA foreign_keys = ON")
    try:
        yield c
        c.commit()
    finally:
        c.close()


def init_db():
    with db() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS arbitros(
            id INTEGER PRIMARY KEY, nombre TEXT NOT NULL,
            nombre_norm TEXT UNIQUE NOT NULL, residencia TEXT);
        CREATE TABLE IF NOT EXISTS usuarios(
            id INTEGER PRIMARY KEY, nombre TEXT NOT NULL,
            nombre_norm TEXT UNIQUE NOT NULL, rol TEXT NOT NULL,
            arbitro_id INTEGER UNIQUE REFERENCES arbitros(id),
            pass_hash TEXT NOT NULL, salt TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS torneos(
            nombre TEXT PRIMARY KEY, paga_pasaje INTEGER NOT NULL DEFAULT 0);
        CREATE TABLE IF NOT EXISTS precios(
            torneo TEXT, categoria TEXT, rol TEXT, valor REAL NOT NULL,
            PRIMARY KEY(torneo, categoria, rol));
        CREATE TABLE IF NOT EXISTS pasajes(
            residencia TEXT, ciudad TEXT, valor REAL NOT NULL,
            PRIMARY KEY(residencia, ciudad));
        CREATE TABLE IF NOT EXISTS servicios(
            id INTEGER PRIMARY KEY, uid TEXT UNIQUE NOT NULL,
            fecha TEXT NOT NULL, hora TEXT, torneo TEXT NOT NULL, ciudad TEXT,
            campo TEXT, categoria TEXT NOT NULL,
            arbitro_id INTEGER NOT NULL REFERENCES arbitros(id),
            partido TEXT, rol TEXT NOT NULL, fase TEXT, cobro TEXT,
            tipo_partido TEXT, archivo TEXT);
        CREATE INDEX IF NOT EXISTS ix_serv ON servicios(arbitro_id, fecha);
        """)


# ------------------------------------------------------------------ usuarios
def _hash(password: str, salt: str) -> str:
    return hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt), 200_000).hex()


def codigo_admin() -> str:
    """Código secreto para crear administradores: variable de entorno o archivo."""
    env = os.environ.get("COARCUN_ADMIN_CODE")
    if env:
        return env
    if not os.path.exists(CODIGO_FILE):
        with open(CODIGO_FILE, "w") as f:
            f.write(secrets.token_urlsafe(8))
    return open(CODIGO_FILE).read().strip()


def crear_usuario(nombre, password, rol, arbitro_id=None):
    if len(password) < 6:
        raise ValueError("La contraseña debe tener al menos 6 caracteres.")
    salt = secrets.token_hex(16)
    try:
        with db() as c:
            c.execute("INSERT INTO usuarios(nombre,nombre_norm,rol,arbitro_id,pass_hash,salt)"
                      " VALUES(?,?,?,?,?,?)",
                      (limpiar(nombre), norm(nombre), rol, arbitro_id, _hash(password, salt), salt))
    except sqlite3.IntegrityError:
        raise ValueError("Ese nombre ya tiene una cuenta.")


def cambiar_password(usuario_id, password):
    if len(password) < 6:
        raise ValueError("La contraseña debe tener al menos 6 caracteres.")
    salt = secrets.token_hex(16)
    with db() as c:
        c.execute("UPDATE usuarios SET pass_hash=?, salt=? WHERE id=?",
                  (_hash(password, salt), salt, usuario_id))


def autenticar(nombre, password):
    with db() as c:
        u = c.execute("SELECT * FROM usuarios WHERE nombre_norm=?", (norm(nombre),)).fetchone()
    if u and hmac.compare_digest(u["pass_hash"], _hash(password, u["salt"])):
        return {"id": u["id"], "nombre": u["nombre"], "rol": u["rol"], "arbitro_id": u["arbitro_id"]}
    return None


def hay_administradores() -> bool:
    with db() as c:
        return c.execute("SELECT 1 FROM usuarios WHERE rol='admin'").fetchone() is not None


def arbitros_sin_cuenta():
    with db() as c:
        return [dict(r) for r in c.execute(
            "SELECT a.* FROM arbitros a LEFT JOIN usuarios u ON u.arbitro_id=a.id "
            "WHERE u.id IS NULL ORDER BY a.nombre")]


def registrar_arbitro(arbitro_id, password, residencia):
    with db() as c:
        a = c.execute("SELECT nombre FROM arbitros WHERE id=?", (arbitro_id,)).fetchone()
    crear_usuario(a["nombre"], password, "arbitro", arbitro_id)  # falla si ya tiene cuenta
    actualizar_residencia(arbitro_id, residencia)


def sugerir_nombres(nombre):
    """Árbitros sin cuenta cuyo nombre se parece al escrito (incluye nombres incompletos)."""
    n = norm(nombre)
    tokens = set(n.split())
    out = []
    for a in arbitros_sin_cuenta():
        cand = set(a["nombre_norm"].split())
        if (tokens and tokens <= cand) or difflib.SequenceMatcher(None, n, a["nombre_norm"]).ratio() >= 0.8:
            out.append(a["nombre"])
    return out[:5]


def registrar_arbitro_por_nombre(nombre, password, residencia, crear_si_no_existe=False):
    """El árbitro escribe su nombre. Si coincide con uno de las designaciones, se enlaza.
    Devuelve ('enlazado' | 'nuevo'). Lanza ValueError con un mensaje claro si no se puede."""
    nombre = limpiar(nombre)
    n = norm(nombre)
    if len(n.split()) < 2:
        raise ValueError("Escribe tu nombre completo (nombres y apellidos).")
    with db() as c:
        a = c.execute("SELECT id FROM arbitros WHERE nombre_norm=?", (n,)).fetchone()
    if a:
        registrar_arbitro(a["id"], password, residencia)  # lanza error si ya tiene cuenta
        return "enlazado"
    if not crear_si_no_existe:
        raise LookupError("No encontré ese nombre en las designaciones.")
    agregar_arbitro(nombre, residencia)
    with db() as c:
        aid = c.execute("SELECT id FROM arbitros WHERE nombre_norm=?", (n,)).fetchone()["id"]
    registrar_arbitro(aid, password, residencia)
    return "nuevo"


def listar_arbitros():
    with db() as c:
        return [dict(r) for r in c.execute(
            "SELECT a.id, a.nombre, a.residencia, u.id AS usuario_id "
            "FROM arbitros a LEFT JOIN usuarios u ON u.arbitro_id=a.id ORDER BY a.nombre")]


def agregar_arbitro(nombre, residencia=None):
    try:
        with db() as c:
            c.execute("INSERT INTO arbitros(nombre,nombre_norm,residencia) VALUES(?,?,?)",
                      (limpiar(nombre), norm(nombre), residencia))
    except sqlite3.IntegrityError:
        raise ValueError("Ese árbitro ya existe.")


def actualizar_residencia(arbitro_id, residencia):
    with db() as c:
        c.execute("UPDATE arbitros SET residencia=? WHERE id=?", (residencia or None, arbitro_id))


# ----------------------------------------------------- importar designaciones
COLUMNAS = {"fecha": "fecha", "torneo": "torneo", "ciudad": "ciudad", "campo": "campo",
            "categoria": "categoria", "hora": "hora", "arbitro": "arbitro", "partido": "partido",
            "servicio": "servicio", "fase": "fase", "cobro": "cobro"}
REQUERIDAS = ["fecha", "torneo", "categoria", "arbitro", "servicio"]


def _leer_tabla(fuente, nombre):
    if nombre.lower().endswith(".csv"):
        try:
            return pd.read_csv(fuente, header=None, dtype=object, encoding="utf-8-sig")
        except UnicodeDecodeError:
            fuente.seek(0)
            return pd.read_csv(fuente, header=None, dtype=object, encoding="latin-1")
    return pd.read_excel(fuente, header=None, dtype=object)


def importar_designaciones(fuente, nombre_archivo="archivo.xlsx") -> dict:
    """Lee el Excel/CSV de designaciones y lo guarda. Devuelve un reporte."""
    raw = _leer_tabla(fuente, nombre_archivo)

    # 1) encontrar la fila de encabezados (el archivo trae un título arriba)
    hdr = next((i for i, r in raw.iterrows()
                if {"fecha", "arbitro"} <= {norm(v) for v in r if pd.notna(v)}), None)
    if hdr is None:
        raise ValueError("No encontré los encabezados 'Fecha' y 'Árbitro' en el archivo.")

    # 2) ubicar columnas por nombre (tolera el typo 'ipo de partid')
    pos = {}
    for j, v in raw.iloc[hdr].items():
        n = norm(v) if pd.notna(v) else ""
        if n in COLUMNAS:
            pos[COLUMNAS[n]] = j
        elif n.startswith(("tipo de", "ipo de")):
            pos["tipo_partido"] = j
    faltan = [k for k in REQUERIDAS if k not in pos]
    if faltan:
        raise ValueError(f"Faltan columnas obligatorias: {', '.join(faltan)}")

    def val(row, k):
        v = row[pos[k]] if k in pos else None
        return None if v is None or (not isinstance(v, (time,)) and pd.isna(v)) else v

    rep = {"filas": 0, "nuevos": 0, "duplicados": 0, "omitidas": [],
           "arbitros_nuevos": [], "posibles_duplicados": [], "roles_raros": {}}
    with db() as c:
        existentes = {r["nombre_norm"]: r["nombre"] for r in c.execute("SELECT nombre_norm,nombre FROM arbitros")}
        for i in range(hdr + 1, len(raw)):
            row = raw.iloc[i]
            if not any(pd.notna(x) for j, x in row.items() if j != pos["fecha"]):
                continue  # fila separadora (solo trae la fecha)
            rep["filas"] += 1
            excel_row = i + 1
            nombre = val(row, "arbitro")
            if nombre is None:
                rep["omitidas"].append((excel_row, "Sin nombre de árbitro"))
                continue
            nombre = limpiar(nombre)
            fecha = parse_fecha(val(row, "fecha"))
            torneo, cat, serv = val(row, "torneo"), val(row, "categoria"), val(row, "servicio")
            problema = ("Fecha no reconocida" if not fecha else "Sin torneo" if torneo is None
                        else "Sin categoría" if cat is None else "Sin servicio (árbitro/asistente)"
                        if serv is None else None)
            if problema:
                rep["omitidas"].append((excel_row, f"{nombre}: {problema}"))
                continue

            n = norm(nombre)
            if n not in existentes:
                parecidos = difflib.get_close_matches(n, list(existentes), n=1, cutoff=0.88)
                c.execute("INSERT INTO arbitros(nombre,nombre_norm) VALUES(?,?)", (nombre, n))
                existentes[n] = nombre
                rep["arbitros_nuevos"].append(nombre)
                if parecidos:
                    rep["posibles_duplicados"].append((nombre, existentes[parecidos[0]]))
            aid = c.execute("SELECT id FROM arbitros WHERE nombre_norm=?", (n,)).fetchone()["id"]

            torneo, cat, rol = normalizar_torneo(torneo), limpiar(cat), normalizar_rol(serv)
            if rol not in ROLES_COBRABLES:
                rep["roles_raros"][rol] = rep["roles_raros"].get(rol, 0) + 1
            hora = parse_hora(val(row, "hora"))
            campo = limpiar(val(row, "campo") or "")
            partido = limpiar(val(row, "partido") or "")
            uid = hashlib.sha1("|".join([fecha, hora, norm(campo), norm(cat), n, rol,
                                         norm(partido)]).encode()).hexdigest()
            c.execute("INSERT OR IGNORE INTO torneos(nombre) VALUES(?)", (torneo,))
            cur = c.execute(
                "INSERT OR IGNORE INTO servicios(uid,fecha,hora,torneo,ciudad,campo,categoria,"
                "arbitro_id,partido,rol,fase,cobro,tipo_partido,archivo) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (uid, fecha, hora, torneo, limpiar(val(row, "ciudad") or ""), campo, cat, aid,
                 partido, rol, limpiar(val(row, "fase") or ""), limpiar(val(row, "cobro") or ""),
                 limpiar(val(row, "tipo_partido") or ""), nombre_archivo))
            rep["nuevos" if cur.rowcount else "duplicados"] += 1
    return rep


# --------------------------------------------------- precios, pasajes, listas
def listar_torneos():
    with db() as c:
        return [r["nombre"] for r in c.execute("SELECT nombre FROM torneos ORDER BY nombre")]


def rango_fechas():
    with db() as c:
        r = c.execute("SELECT MIN(fecha) a, MAX(fecha) b FROM servicios").fetchone()
    return (date.fromisoformat(r["a"]), date.fromisoformat(r["b"])) if r["a"] else None


def tabla_torneos():
    with db() as c:
        return pd.read_sql_query("SELECT nombre AS Torneo, paga_pasaje AS 'Paga pasaje' "
                                 "FROM torneos ORDER BY nombre", c).astype({"Paga pasaje": bool})


def guardar_torneos(df):
    with db() as c:
        for _, r in df.iterrows():
            c.execute("UPDATE torneos SET paga_pasaje=? WHERE nombre=?", (int(r["Paga pasaje"]), r["Torneo"]))


def tabla_precios():
    """Todas las combinaciones torneo/categoría/rol que existen en servicios + su precio."""
    with db() as c:
        df = pd.read_sql_query(
            "SELECT s.torneo AS Torneo, s.categoria AS Categoría, s.rol AS Rol, p.valor AS Valor "
            "FROM (SELECT DISTINCT torneo,categoria,rol FROM servicios WHERE rol IN (?,?)) s "
            "LEFT JOIN precios p ON p.torneo=s.torneo AND p.categoria=s.categoria AND p.rol=s.rol "
            "ORDER BY s.torneo, s.categoria, s.rol", c, params=ROLES_COBRABLES)
    return df


def guardar_precios(df):
    with db() as c:
        for _, r in df.iterrows():
            if pd.notna(r["Valor"]):
                c.execute("INSERT OR REPLACE INTO precios VALUES(?,?,?,?)",
                          (r["Torneo"], r["Categoría"], r["Rol"], float(r["Valor"])))
            else:
                c.execute("DELETE FROM precios WHERE torneo=? AND categoria=? AND rol=?",
                          (r["Torneo"], r["Categoría"], r["Rol"]))


def tabla_pasajes(residencia):
    with db() as c:
        ciudades = [r["ciudad"] for r in c.execute(
            "SELECT DISTINCT ciudad FROM servicios WHERE ciudad<>'' ORDER BY ciudad")]
        actual = {r["ciudad"]: r["valor"] for r in c.execute(
            "SELECT ciudad,valor FROM pasajes WHERE residencia=?", (residencia,))}
    return pd.DataFrame({"Ciudad del partido": ciudades,
                         "Pasaje": [actual.get(x.upper()) for x in ciudades]})


def guardar_pasajes(residencia, df):
    with db() as c:
        for _, r in df.iterrows():
            ciu = str(r["Ciudad del partido"]).upper()
            if pd.notna(r["Pasaje"]):
                c.execute("INSERT OR REPLACE INTO pasajes VALUES(?,?,?)", (residencia, ciu, float(r["Pasaje"])))
            else:
                c.execute("DELETE FROM pasajes WHERE residencia=? AND ciudad=?", (residencia, ciu))


# -------------------------------------------------------------------- cuenta
def calcular_cuenta(arbitro_id, filtros) -> dict:
    """filtros: lista de {'torneo': str, 'desde': date, 'hasta': date}."""
    with db() as c:
        arb = c.execute("SELECT * FROM arbitros WHERE id=?", (arbitro_id,)).fetchone()
        frames = []
        for f in filtros:
            q = "SELECT * FROM servicios WHERE arbitro_id=? AND fecha BETWEEN ? AND ?"
            p = [arbitro_id, f["desde"].isoformat(), f["hasta"].isoformat()]
            if f.get("torneo") and f["torneo"] != TODOS:
                q += " AND torneo=?"
                p.append(f["torneo"])
            frames.append(pd.read_sql_query(q, c, params=p))
        precios = {(r["torneo"], r["categoria"], r["rol"]): r["valor"] for r in c.execute("SELECT * FROM precios")}
        paga = {r["nombre"]: bool(r["paga_pasaje"]) for r in c.execute("SELECT * FROM torneos")}
        pasajes = {(r["residencia"], r["ciudad"]): r["valor"] for r in c.execute("SELECT * FROM pasajes")}

    df = pd.concat(frames, ignore_index=True).drop_duplicates("id") if frames else pd.DataFrame()
    res = {"arbitro": arb["nombre"], "residencia": arb["residencia"], "filas": [],
           "total": 0.0, "sin_precio": [], "otros": {}, "sin_pasaje": [], "sin_residencia": False}
    if df.empty:
        return res

    otros = df[~df["rol"].isin(ROLES_COBRABLES)]
    res["otros"] = otros["rol"].value_counts().to_dict()
    df = df[df["rol"].isin(ROLES_COBRABLES)]

    for (torneo, fecha, cat), g in df.sort_values(["fecha", "torneo"]).groupby(["torneo", "fecha", "categoria"], sort=False):
        n_as = int((g["rol"] == ROL_ASISTENTE).sum())
        n_c = int((g["rol"] == ROL_CENTRAL).sum())
        v_as = precios.get((torneo, cat, ROL_ASISTENTE))
        v_c = precios.get((torneo, cat, ROL_CENTRAL))
        aviso = []
        for n, v, rol in ((n_as, v_as, ROL_ASISTENTE), (n_c, v_c, ROL_CENTRAL)):
            if n and v is None:
                aviso.append(f"sin precio de {rol}")
                if (torneo, cat, rol) not in res["sin_precio"]:
                    res["sin_precio"].append((torneo, cat, rol))
        subtotal = n_as * (v_as or 0) + n_c * (v_c or 0)

        pasaje = 0.0
        if paga.get(torneo):
            if not arb["residencia"]:
                res["sin_residencia"] = True
            else:
                ciudad = g["ciudad"].mode().iloc[0].upper() if g["ciudad"].notna().any() else ""
                pasaje = pasajes.get((arb["residencia"], ciudad))
                if pasaje is None:
                    pasaje = 0.0
                    aviso.append(f"sin pasaje {arb['residencia']}→{ciudad.title()}")
                    if (arb["residencia"], ciudad) not in res["sin_pasaje"]:
                        res["sin_pasaje"].append((arb["residencia"], ciudad))
        res["filas"].append({"Torneo": torneo, "Fecha": fecha, "Categ.": cat,
                             "Cant. As.": n_as, "Valor unit. Asistencias": v_as or 0,
                             "Cant. C.": n_c, "Valor unit. Centros": v_c or 0,
                             "Subtotal": subtotal, "Pasaje": pasaje,
                             "Valor Total": subtotal + pasaje, "Aviso": "; ".join(aviso)})
        res["total"] += subtotal + pasaje
    return res


def exportar_xlsx(res: dict) -> bytes:
    """Excel con el mismo estilo de la tabla de ejemplo (encabezado azul)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Cuenta"
    azul = PatternFill("solid", fgColor="0070C0")
    blanco = Font(name="Arial", bold=True, color="FFFFFF")
    borde = Border(*(Side(style="thin", color="999999"),) * 4)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws["A1"] = f"Cuenta de cobro - {res['arbitro']}"
    ws["A1"].font = Font(name="Arial", bold=True, size=13)
    ws.append([])
    enc = ["Torneo", "Fecha", "Categ.", "Cant. As.", "Valor unitario Asistencias",
           "Cant. C.", "Valor unitario Centros", "Valor Total"]
    ws.append(enc)
    for cell in ws[3]:
        cell.fill, cell.font, cell.alignment, cell.border = azul, blanco, centro, borde
    for f in res["filas"]:
        total = fmt(f["Subtotal"]) + (f"\nPasaje: {fmt(f['Pasaje'])}" if f["Pasaje"] else "")
        ws.append([f["Torneo"], fecha_bonita(f["Fecha"]), f["Categ."], f["Cant. As."],
                   f["Valor unit. Asistencias"], f["Cant. C."], f["Valor unit. Centros"], total])
        for cell in ws[ws.max_row]:
            cell.border, cell.alignment, cell.font = borde, centro, Font(name="Arial")
        for col in ("E", "G"):
            ws[f"{col}{ws.max_row}"].number_format = "#,##0"
    ws.append(["Total General", "", "", "", "", "", "", res["total"]])
    r = ws.max_row
    ws.merge_cells(f"A{r}:G{r}")
    for cell in ws[r]:
        cell.fill, cell.font, cell.alignment, cell.border = azul, blanco, centro, borde
    ws[f"H{r}"].number_format = "#,##0"
    for col, w in zip("ABCDEFGH", (22, 16, 10, 10, 18, 10, 18, 20)):
        ws.column_dimensions[col].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
